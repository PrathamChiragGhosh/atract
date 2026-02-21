#!/usr/bin/env python3
"""
Enhanced Python Document Service
Handles PDF, DOC, DOCX, TXT, Excel file parsing with Spire.Doc support
Integrates with existing filtering pipeline
"""

import json
import sys
import os
import re
from pathlib import Path
from typing import List, Dict, Any, Optional
from datetime import datetime

# Spire libraries
try:
    from spire.doc import Document as SpireWordDocument
    from spire.xls import Workbook as SpireWorkbook
    HAS_SPIRE_DOC = True
except ImportError:
    HAS_SPIRE_DOC = False

try:
    from spire.pdf import PdfDocument, PdfTextExtractOptions, PdfTextExtractor
    HAS_SPIRE_PDF = True
except ImportError:
    HAS_SPIRE_PDF = False

# Fallback libraries
try:
    import PyPDF2
    HAS_PYPDF2 = True
except ImportError:
    HAS_PYPDF2 = False

try:
    import pdfplumber
    HAS_PDFPLUMBER = True
except ImportError:
    HAS_PDFPLUMBER = False

try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import pandas as pd
    HAS_PANDAS = True
except ImportError:
    HAS_PANDAS = False

# Import existing filtering modules
try:
    from parse_and_filter_files import process_files_and_filter
    HAS_FILTERING = True
except ImportError:
    HAS_FILTERING = False
    print("Warning: Filtering module not found", file=sys.stderr)


def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from PDF using Spire or fallback"""
    text = ""
    
    # Try Spire PDF first
    if HAS_SPIRE_PDF:
        try:
            pdf_doc = PdfDocument()
            pdf_doc.LoadFromFile(file_path)
            extractOptions = PdfTextExtractOptions()
            extractOptions.IsExtractAllText = True
            
            for i in range(pdf_doc.Pages.Count):
                page = pdf_doc.Pages.get_Item(i)
                text += PdfTextExtractor(page).ExtractText(extractOptions) + "\n"
            
            pdf_doc.Close()
            if text.strip():
                return text.strip()
        except Exception as e:
            print(f"Spire PDF extraction failed: {e}, trying fallback...", file=sys.stderr)
    
    # Fallback to pdfplumber
    if HAS_PDFPLUMBER:
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            if text.strip():
                return text.strip()
        except Exception as e:
            print(f"pdfplumber failed: {e}, trying PyPDF2...", file=sys.stderr)
    
    # Fallback to PyPDF2
    if HAS_PYPDF2:
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            return text.strip()
        except Exception as e:
            raise Exception(f"Failed to extract text from PDF: {e}")
    
    raise Exception("No PDF parsing library available")


def extract_text_from_word(file_path: str) -> str:
    """Extract text from Word document using Spire or fallback"""
    # Try Spire first
    if HAS_SPIRE_DOC:
        try:
            doc = SpireWordDocument()
            doc.LoadFromFile(file_path)
            text = doc.GetText()
            doc.Close()
            return text.strip()
        except Exception as e:
            print(f"Spire Word extraction failed: {e}, trying fallback...", file=sys.stderr)
    
    # Fallback to python-docx
    if HAS_DOCX:
        try:
            doc = Document(file_path)
            text_parts = []
            
            # Extract paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            
            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        text_parts.append(" | ".join(row_text))
            
            return "\n".join(text_parts).strip()
        except Exception as e:
            raise Exception(f"Failed to extract text from Word document: {e}")
    
    raise Exception("No Word parsing library available. Install spire-doc or python-docx")


def extract_data_from_excel(file_path: str) -> List[Dict[str, Any]]:
    """Extract candidate data from Excel file using Spire or fallback"""
    candidates = []
    
    # Try Spire XLS first
    if HAS_SPIRE_DOC:
        try:
            workbook = SpireWorkbook()
            workbook.LoadFromFile(file_path)
            sheet = workbook.Worksheets[0]
            
            # Find header row
            headers = []
            header_row = 1
            max_cols = sheet.LastColumn
            
            # Read header row
            for col in range(1, max_cols + 1):
                cell_value = sheet.Range[header_row, col].Value
                headers.append(str(cell_value) if cell_value else f'Column{col}')
            
            # Read data rows
            max_rows = sheet.LastRow
            for row in range(header_row + 1, max_rows + 1):
                candidate = {}
                for col_idx, header in enumerate(headers, 1):
                    cell_value = sheet.Range[row, col_idx].Value
                    candidate[header] = str(cell_value) if cell_value else ''
                
                if any(candidate.values()):  # Skip empty rows
                    candidates.append(candidate)
            
            workbook.Dispose()
            return candidates
        except Exception as e:
            print(f"Spire Excel extraction failed: {e}, trying fallback...", file=sys.stderr)
    
    # Fallback to openpyxl/pandas
    if HAS_PANDAS:
        try:
            if file_path.endswith('.xlsx'):
                df = pd.read_excel(file_path, engine='openpyxl')
            else:
                df = pd.read_excel(file_path, engine='xlrd')
            
            candidates = df.to_dict('records')
            for candidate in candidates:
                for key, value in candidate.items():
                    candidate[key] = str(value) if pd.notna(value) else ''
            
            return candidates
        except Exception as e:
            raise Exception(f"Failed to extract data from Excel: {e}")
    
    raise Exception("No Excel parsing library available")


def extract_text_from_txt(file_path: str) -> str:
    """Extract text from plain text file"""
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
            return file.read().strip()
    except Exception as e:
        raise Exception(f"Failed to read text file: {e}")


def parse_file(file_path: str) -> Dict[str, Any]:
    """Parse a file and extract content based on file type"""
    file_path_obj = Path(file_path)
    
    if not file_path_obj.exists():
        raise Exception(f"File not found: {file_path}")
    
    file_ext = file_path_obj.suffix.lower()
    file_name = file_path_obj.name
    
    result = {
        'fileName': file_name,
        'fileType': file_ext.lstrip('.'),
        'candidates': [],
        'text': ''
    }
    
    try:
        if file_ext == '.pdf':
            result['text'] = extract_text_from_pdf(str(file_path_obj))
        elif file_ext in ['.docx', '.doc']:
            if file_ext == '.doc' and not HAS_SPIRE_DOC:
                raise Exception(".doc files require Spire.Doc. Please convert to .docx or install spire-doc")
            result['text'] = extract_text_from_word(str(file_path_obj))
        elif file_ext in ['.xlsx', '.xls']:
            result['candidates'] = extract_data_from_excel(str(file_path_obj))
        elif file_ext == '.txt':
            result['text'] = extract_text_from_txt(str(file_path_obj))
        else:
            raise Exception(f"Unsupported file type: {file_ext}")
        
        return result
    except Exception as e:
        raise Exception(f"Error parsing file {file_name}: {e}")


def convert_text_to_candidate(text: str, file_name: str) -> Dict[str, Any]:
    """Convert extracted text to candidate format"""
    candidate = {
        'name': '',
        'email': '',
        'mobile': '',
        'age': '',
        'gender': '',
        'experience': '',
        'skills': '',
        'location': '',
        'qualification': '',
        'currentRole': '',
        'noticePeriod': '',
        'salary': '',
        'extractedText': text,
        'sourceFile': file_name
    }
    
    # Extract email
    email_match = re.search(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', text)
    if email_match:
        candidate['email'] = email_match.group(0)
    
    # Extract phone/mobile
    phone_patterns = [
        r'\b\d{10}\b',
        r'\b\d{3}[-.\s]?\d{3}[-.\s]?\d{4}\b',
        r'\b\+?\d{1,3}[-.\s]?\d{1,4}[-.\s]?\d{1,4}[-.\s]?\d{1,9}\b'
    ]
    for pattern in phone_patterns:
        phone_match = re.search(pattern, text)
        if phone_match:
            candidate['mobile'] = phone_match.group(0)
            break
    
    # Extract name (first line or before email)
    lines = text.split('\n')[:5]
    for line in lines:
        line = line.strip()
        if line and '@' not in line and not line.isdigit():
            candidate['name'] = line
            break
    
    if not candidate['name']:
        candidate['name'] = Path(file_name).stem
    
    return candidate


def process_and_filter_files(file_paths: List[str], job: Dict[str, Any], 
                             filter_percentage: float = 50.0,
                             use_llm: bool = True) -> Dict[str, Any]:
    """Process files and apply filtering"""
    if HAS_FILTERING:
        return process_files_and_filter(
            file_paths,
            job,
            filter_percentage,
            use_llm
        )
    else:
        # Basic processing without filtering
        all_candidates = []
        for file_path in file_paths:
            try:
                parse_result = parse_file(file_path)
                if parse_result['fileType'] in ['xlsx', 'xls']:
                    for excel_candidate in parse_result['candidates']:
                        candidate = {
                            **excel_candidate,
                            'sourceFile': parse_result['fileName'],
                            'fileType': parse_result['fileType']
                        }
                        all_candidates.append(candidate)
                else:
                    candidate = convert_text_to_candidate(
                        parse_result['text'],
                        parse_result['fileName']
                    )
                    candidate['fileType'] = parse_result['fileType']
                    all_candidates.append(candidate)
            except Exception as e:
                print(f"Error processing {file_path}: {e}", file=sys.stderr)
        
        return {
            'success': True,
            'totalCandidates': len(all_candidates),
            'candidates': all_candidates,
            'message': 'Filtering module not available, returning all candidates'
        }


def main():
    """Main function - expects JSON input from stdin"""
    try:
        input_data = json.load(sys.stdin)
        
        action = input_data.get('action', 'parse')
        
        if action == 'parse':
            # Single file parsing
            file_path = input_data.get('filePath')
            if not file_path:
                raise ValueError("filePath required for parse action")
            
            result = parse_file(file_path)
            print(json.dumps({
                'success': True,
                'result': result
            }))
        
        elif action == 'process_and_filter':
            # Process multiple files and filter
            file_paths = input_data.get('filePaths', [])
            job = input_data.get('job', {})
            filter_percentage = float(input_data.get('filterPercentage', 50.0))
            use_llm = input_data.get('useLLM', True)
            
            if not file_paths:
                raise ValueError("filePaths required")
            if not job:
                raise ValueError("job requirements required")
            
            result = process_and_filter_files(
                file_paths,
                job,
                filter_percentage,
                use_llm
            )
            print(json.dumps(result))
        
        else:
            raise ValueError(f"Unknown action: {action}")
    
    except json.JSONDecodeError as e:
        print(json.dumps({
            'success': False,
            'error': f'Invalid JSON input: {str(e)}'
        }), file=sys.stderr)
        sys.exit(1)
    except Exception as e:
        print(json.dumps({
            'success': False,
            'error': f'Processing error: {str(e)}'
        }), file=sys.stderr)
        import traceback
        traceback.print_exc(file=sys.stderr)
        sys.exit(1)


if __name__ == '__main__':
    main()

