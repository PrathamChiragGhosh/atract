#!/usr/bin/env python3
"""
Complete Python-based File Parser and Filter System
Reads files (PDF, Word, Excel, TXT) using Python libraries
Then applies two-stage filtering: Vector Similarity + LLM Evaluation
"""

import json
import sys
import os
import re
from typing import List, Dict, Any, Optional
from datetime import datetime
from pathlib import Path

# File parsing libraries
try:
    import PyPDF2
    HAS_PYPDF2 = True
except ImportError:
    HAS_PYPDF2 = False

try:
    import pdfplumber
    HAS_PDFPLUMBER = True
except ImportError:
    HAS_PDFPLUMBER = False

try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import pandas as pd
    HAS_PANDAS = True
except ImportError:
    HAS_PANDAS = False

try:
    import xlrd
    HAS_XLRD = True
except ImportError:
    HAS_XLRD = False

# Import filtering functions from filter_candidates.py
try:
    from filter_candidates import (
        filter_candidates_vector,
        build_resume_text,
        build_job_text
    )
    HAS_FILTER_MODULE = True
except ImportError:
    HAS_FILTER_MODULE = False
    print("Warning: filter_candidates module not found", file=sys.stderr)

# Import LLM evaluation
try:
    from evaluate_candidate_llm import evaluate_with_llm
    HAS_LLM_MODULE = True
except ImportError:
    HAS_LLM_MODULE = False
    print("Warning: LLM evaluation module not found", file=sys.stderr)

def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from PDF file"""
    text = ""
    
    # Try pdfplumber first (better text extraction)
    if HAS_PDFPLUMBER:
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            if text.strip():
                return text.strip()
        except Exception as e:
            print(f"pdfplumber failed: {e}, trying PyPDF2...", file=sys.stderr)
    
    # Fallback to PyPDF2
    if HAS_PYPDF2:
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            return text.strip()
        except Exception as e:
            raise Exception(f"Failed to extract text from PDF: {e}")
    
    raise Exception("No PDF parsing library available. Install pdfplumber or PyPDF2")

def extract_text_from_word(file_path: str) -> str:
    """Extract text from Word document (.docx)"""
    if not HAS_DOCX:
        raise Exception("python-docx not installed. Install with: pip install python-docx")
    
    try:
        doc = Document(file_path)
        text_parts = []
        
        # Extract text from paragraphs
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)
        
        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    text_parts.append(" | ".join(row_text))
        
        return "\n".join(text_parts).strip()
    except Exception as e:
        raise Exception(f"Failed to extract text from Word document: {e}")

def extract_data_from_excel(file_path: str) -> List[Dict[str, Any]]:
    """Extract candidate data from Excel file"""
    candidates = []
    
    # Try openpyxl first (for .xlsx)
    if HAS_OPENPYXL and file_path.endswith('.xlsx'):
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            sheet = workbook.active
            
            # Find header row
            headers = []
            header_row = None
            for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=True), 1):
                row_values = [str(cell).lower().strip() if cell else '' for cell in row]
                # Check if this row contains header keywords
                header_keywords = ['name', 'email', 'mobile', 'phone', 'age', 'gender', 
                                 'location', 'experience', 'skills', 'qualification']
                matches = sum(1 for val in row_values if any(keyword in val for keyword in header_keywords))
                if matches >= 3:
                    headers = [str(cell) if cell else f'Column{i+1}' for i, cell in enumerate(row)]
                    header_row = row_idx
                    break
            
            if not headers:
                headers = [str(cell) if cell else f'Column{i+1}' 
                          for i, cell in enumerate(next(sheet.iter_rows(values_only=True)))]
                header_row = 1
            
            # Extract data rows
            for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                if not any(cell for cell in row):  # Skip empty rows
                    continue
                
                candidate = {}
                for idx, cell_value in enumerate(row):
                    if idx < len(headers):
                        candidate[headers[idx]] = str(cell_value) if cell_value else ''
                
                if candidate:
                    candidates.append(candidate)
            
            return candidates
        except Exception as e:
            print(f"openpyxl failed: {e}, trying pandas...", file=sys.stderr)
    
    # Fallback to pandas (works for both .xlsx and .xls)
    if HAS_PANDAS:
        try:
            # Try .xlsx first
            if file_path.endswith('.xlsx'):
                df = pd.read_excel(file_path, engine='openpyxl')
            else:
                df = pd.read_excel(file_path, engine='xlrd')
            
            # Convert to list of dictionaries
            candidates = df.to_dict('records')
            # Convert all values to strings
            for candidate in candidates:
                for key, value in candidate.items():
                    candidate[key] = str(value) if pd.notna(value) else ''
            
            return candidates
        except Exception as e:
            raise Exception(f"Failed to extract data from Excel: {e}")
    
    raise Exception("No Excel parsing library available. Install openpyxl or pandas")

def extract_text_from_txt(file_path: str) -> str:
    """Extract text from plain text file"""
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
            return file.read().strip()
    except Exception as e:
        raise Exception(f"Failed to read text file: {e}")

def parse_file(file_path: str) -> Dict[str, Any]:
    """
    Parse a file and extract content based on file type
    
    Returns:
        {
            'fileType': 'pdf'|'docx'|'xlsx'|'txt',
            'candidates': [...],  # For Excel files
            'text': '...',  # For PDF/Word/TXT files
            'fileName': '...'
        }
    """
    file_path = Path(file_path)
    
    if not file_path.exists():
        raise Exception(f"File not found: {file_path}")
    
    file_ext = file_path.suffix.lower()
    file_name = file_path.name
    
    result = {
        'fileName': file_name,
        'fileType': file_ext.lstrip('.'),
        'candidates': [],
        'text': ''
    }
    
    try:
        if file_ext == '.pdf':
            result['text'] = extract_text_from_pdf(str(file_path))
        elif file_ext in ['.docx', '.doc']:
            if file_ext == '.doc':
                raise Exception(".doc files are not supported. Please convert to .docx")
            result['text'] = extract_text_from_word(str(file_path))
        elif file_ext in ['.xlsx', '.xls']:
            result['candidates'] = extract_data_from_excel(str(file_path))
        elif file_ext == '.txt':
            result['text'] = extract_text_from_txt(str(file_path))
        else:
            raise Exception(f"Unsupported file type: {file_ext}")
        
        return result
    except Exception as e:
        raise Exception(f"Error parsing file {file_name}: {e}")

def convert_text_to_candidate(text: str, file_name: str) -> Dict[str, Any]:
    """Convert extracted text to candidate format"""
    # Basic extraction of common fields
    candidate = {
        'name': '',
        'email': '',
        'mobile': '',
        'age': '',
        'gender': '',
        'experience': '',
        'skills': '',
        'location': '',
        'qualification': '',
        'currentRole': '',
        'noticePeriod': '',
        'salary': '',
        'extractedText': text,
        'sourceFile': file_name
    }
    
    # Extract email
    email_match = re.search(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', text)
    if email_match:
        candidate['email'] = email_match.group(0)
    
    # Extract phone/mobile
    phone_patterns = [
        r'\b\d{10}\b',  # 10 digits
        r'\b\d{3}[-.\s]?\d{3}[-.\s]?\d{4}\b',  # US format
        r'\b\+?\d{1,3}[-.\s]?\d{1,4}[-.\s]?\d{1,4}[-.\s]?\d{1,9}\b'  # International
    ]
    for pattern in phone_patterns:
        phone_match = re.search(pattern, text)
        if phone_match:
            candidate['mobile'] = phone_match.group(0)
            break
    
    # Extract name (first line or before email)
    lines = text.split('\n')[:5]
    for line in lines:
        line = line.strip()
        if line and '@' not in line and not line.isdigit():
            candidate['name'] = line
            break
    
    # Use filename as name if not found
    if not candidate['name']:
        candidate['name'] = Path(file_name).stem
    
    return candidate

def process_files_and_filter(file_paths: List[str], job: Dict[str, Any], 
                             python_filter_percentage: float = 50.0,
                             use_llm: bool = True) -> Dict[str, Any]:
    """
    Process multiple files, extract content, and apply two-stage filtering
    
    Args:
        file_paths: List of file paths to process
        job: Job requirements dictionary
        python_filter_percentage: Percentage for Stage 1 filtering (0-100)
        use_llm: Whether to use Stage 2 LLM evaluation
    
    Returns:
        Filtered candidates with match scores
    """
    all_candidates = []
    processed_files = []
    
    print(f"Processing {len(file_paths)} file(s)...", file=sys.stderr)
    
    # Step 1: Parse all files and extract candidates
    for file_path in file_paths:
        try:
            print(f"Parsing: {file_path}", file=sys.stderr)
            parse_result = parse_file(file_path)
            
            if parse_result['fileType'] in ['xlsx', 'xls']:
                # Excel file: multiple candidates
                for excel_candidate in parse_result['candidates']:
                    candidate = {
                        **excel_candidate,
                        'sourceFile': parse_result['fileName'],
                        'fileType': parse_result['fileType']
                    }
                    all_candidates.append(candidate)
                processed_files.append({
                    'file': parse_result['fileName'],
                    'type': parse_result['fileType'],
                    'candidates': len(parse_result['candidates']),
                    'status': 'success'
                })
            else:
                # PDF/Word/TXT: single candidate per file
                candidate = convert_text_to_candidate(
                    parse_result['text'],
                    parse_result['fileName']
                )
                candidate['fileType'] = parse_result['fileType']
                all_candidates.append(candidate)
                processed_files.append({
                    'file': parse_result['fileName'],
                    'type': parse_result['fileType'],
                    'candidates': 1,
                    'status': 'success'
                })
        except Exception as e:
            print(f"Error processing {file_path}: {e}", file=sys.stderr)
            processed_files.append({
                'file': Path(file_path).name,
                'type': Path(file_path).suffix,
                'candidates': 0,
                'status': 'error',
                'error': str(e)
            })
    
    if not all_candidates:
        return {
            'success': False,
            'error': 'No candidates extracted from files',
            'processedFiles': processed_files
        }
    
    print(f"Extracted {len(all_candidates)} candidate(s) from files", file=sys.stderr)
    
    # Step 2: Stage 1 - Vector-based filtering
    if not HAS_FILTER_MODULE:
        return {
            'success': False,
            'error': 'Filter module not available',
            'candidates': all_candidates,
            'processedFiles': processed_files
        }
    
    print(f"Stage 1: Vector similarity filtering ({python_filter_percentage}%)...", file=sys.stderr)
    filtered_candidates = filter_candidates_vector(
        all_candidates,
        job,
        python_filter_percentage
    )
    
    print(f"Stage 1 completed: {len(filtered_candidates)} candidate(s) after filtering", file=sys.stderr)
    
    # Step 3: Stage 2 - LLM evaluation (optional)
    if use_llm and HAS_LLM_MODULE:
        print(f"Stage 2: LLM evaluation...", file=sys.stderr)
        evaluated_candidates = []
        
        for candidate in filtered_candidates:
            try:
                evaluation = evaluate_with_llm(candidate, job)
                evaluated_candidate = {
                    **candidate,
                    'matchScore': evaluation['matchScore'],
                    'freshnessScore': evaluation['freshnessScore'],
                    'keyMatchingSkills': evaluation['keyMatchingSkills'],
                    'explanation': evaluation['explanation'],
                    'evaluationDetails': {
                        'skillMatch': evaluation['skillMatch'],
                        'roleRelevance': evaluation['roleRelevance'],
                        'careerProgression': evaluation['careerProgression'],
                        'overallFit': evaluation['overallFit'],
                        'shouldReject': evaluation['shouldReject']
                    }
                }
                
                # Only include if not rejected
                if not evaluation.get('shouldReject', False):
                    evaluated_candidates.append(evaluated_candidate)
            except Exception as e:
                print(f"LLM evaluation error for {candidate.get('name', 'unknown')}: {e}", file=sys.stderr)
                evaluated_candidates.append(candidate)
        
        # Re-sort by match score
        evaluated_candidates.sort(key=lambda x: x.get('matchScore', 0), reverse=True)
        final_candidates = evaluated_candidates
    else:
        final_candidates = filtered_candidates
    
    return {
        'success': True,
        'totalFiles': len(file_paths),
        'totalCandidates': len(all_candidates),
        'filteredCandidates': len(final_candidates),
        'pythonFilterPercentage': python_filter_percentage,
        'usedLLM': use_llm and HAS_LLM_MODULE,
        'candidates': final_candidates,
        'processedFiles': processed_files
    }

def main():
    """Main function"""
    try:
        input_data = json.load(sys.stdin)
        
        file_paths = input_data.get('filePaths', [])
        job = input_data.get('job', {})
        python_filter_percentage = float(input_data.get('pythonFilterPercentage', 50.0))
        use_llm = input_data.get('useLLM', True)
        
        if not file_paths:
            print(json.dumps({
                'success': False,
                'error': 'No file paths provided'
            }), file=sys.stderr)
            sys.exit(1)
        
        if not job:
            print(json.dumps({
                'success': False,
                'error': 'No job requirements provided'
            }), file=sys.stderr)
            sys.exit(1)
        
        # Process files and filter
        result = process_files_and_filter(
            file_paths,
            job,
            python_filter_percentage,
            use_llm
        )
        
        print(json.dumps(result))
        
    except json.JSONDecodeError as e:
        print(json.dumps({
            'success': False,
            'error': f'Invalid JSON input: {str(e)}'
        }), file=sys.stderr)
        sys.exit(1)
    except Exception as e:
        print(json.dumps({
            'success': False,
            'error': f'Processing error: {str(e)}'
        }), file=sys.stderr)
        import traceback
        traceback.print_exc(file=sys.stderr)
        sys.exit(1)

if __name__ == '__main__':
    main()

